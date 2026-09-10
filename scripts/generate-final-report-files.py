import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, color_hex):
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def generate_docx(output_path):
    doc = docx.Document()
    
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Segoe UI'
    normal_style.font.size = Pt(10)
    normal_style.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    # Title
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run('SAKUKILAT — LAPORAN RESMI\nFINAL CORRECTION, UX UPGRADE & CATEGORY ANALYTICS')
    title_run.font.name = 'Arial'
    title_run.font.size = Pt(17)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_p.add_run('Dokumentasi Resmi Penyelesaian Remediasi, Navigasi Back Stack, Input Rupiah, UI Catat Transaksi, dan Jejak Kategori Setahun')
    sub_run.font.size = Pt(10)
    sub_run.font.italic = True
    sub_run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Metadata Box
    meta_table = doc.add_table(rows=6, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Tanggal Pelaksanaan", "11 September 2026"),
        ("Branch Kerja", "feat/navigation-entry-category-responsive"),
        ("Base Commit & Origin", "f01f76e (origin/main)"),
        ("Current Local HEAD", "1bae659 (8 commits ahead of origin/main)"),
        ("Status Eksekusi", "SELESAI PENUH (8/8 Commits, 8 Suited Tests PASS, 2 Builds PASS)"),
        ("Repository", "C:\\Users\\HYPE AMD\\Projects\\SakuKilat")
    ]
    for idx, (label, val) in enumerate(meta_data):
        row = meta_table.rows[idx]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(2.2)
        c1.width = Inches(4.8)
        set_cell_background(c0, "F1F5F9")
        set_cell_background(c1, "F8FAFC")
        set_cell_margins(c0, 70, 70, 110, 110)
        set_cell_margins(c1, 70, 70, 110, 110)
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(label)
        r0.bold = True
        r0.font.size = Pt(9.5)
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(val)
        r1.font.size = Pt(9.5)
        if label == "Status Eksekusi":
            r1.bold = True
            r1.font.color.rgb = RGBColor(0x04, 0x78, 0x57)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Section 1: Ringkasan Eksekutif
    h1 = doc.add_heading('1. Ringkasan Eksekutif', level=1)
    h1.style.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    p = doc.add_paragraph(
        'Seluruh rangkaian tugas perbaikan tingkat lanjut dan penyempurnaan pengalaman pengguna pada aplikasi SakuKilat '
        'telah berhasil diselesaikan 100% tanpa menyimpang dari batasan teknis (tanpa dependensi eksternal yang membengkak, '
        'tanpa fitur rekonsiliasi saldo/cloud-sync di luar cakupan, serta dengan menjaga kepatuhan data offline lokal).\n\n'
        'Penyelesaian mencakup 7 pilar utama:\n'
        '1. Keselamatan Data & Storage: Proteksi skema masa depan, parser CSV RFC 4180 ketat, dedup import engine, dan pelaporan rollback yang jujur.\n'
        '2. Navigasi Sublayer: Implementasi LIFO back-stack terintegrasi hardware back button Android (Capacitor) dan browser popstate.\n'
        '3. Standardisasi Live Rupiah: Komponen input nominal seragam dengan format ribuan live, prefix protektif, cursor tracking, dan paste shortcuts.\n'
        '4. Overhaul UI Pencatatan: Segmented type selector ergonomis (≥44px), Amount Hero display, pengurutan kategori dinamis berbasis frekuensi, dan tombol submit kontekstual.\n'
        '5. Fitur Baru Jejak Kategori Setahun: Agregasi 12 bulan kalender, visualisasi batang interaktif Recharts, 5 metrik performa finansial, dan drilldown transaksi bulanan.\n'
        '6. Optimasi Responsivitas: Proteksi overflow horizontal 320px+, utilities scrollbar halus, dan padding safe-area.\n'
        '7. Verifikasi Produksi: 7 regression test suites (79 unit tests) + 1 security scanner lulus 100%, kompilasi Next.js Web dan Mobile Export sukses tanpa error TypeScript.'
    )

    # Section 2: Commit History
    h2 = doc.add_heading('2. Riwayat Commit Terverifikasi', level=1)
    h2.style.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    
    commits = [
        ("4e1eed3", "fix(storage)", "Prevent destructive incompatible reset and preserve future keys"),
        ("38bfc6d", "fix(import)", "Strict CSV parser, dedup engine, and honest rollback reporting"),
        ("2497937", "test(navigation)", "Define nested back-stack behavior"),
        ("3c336bd", "fix(navigation)", "Return sublayers to their parent via LIFO back-stack"),
        ("e0546a0", "fix(amount)", "Apply live Rupiah formatting consistently across forms"),
        ("57e9bf7", "feat(entry)", "Improve expense and income entry UI"),
        ("fae2f13", "feat(category-year)", "Add yearly category explorer"),
        ("1bae659", "fix(responsive)", "Optimize supported viewport sizes")
    ]
    c_table = doc.add_table(rows=len(commits)+1, cols=3)
    c_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = c_table.rows[0].cells
    hdr_cells[0].width = Inches(1.2)
    hdr_cells[1].width = Inches(1.8)
    hdr_cells[2].width = Inches(4.0)
    for i, title in enumerate(["Commit Hash", "Scope / Type", "Pesan Commit"]):
        set_cell_background(hdr_cells[i], "1E293B")
        set_cell_margins(hdr_cells[i], 80, 80, 100, 100)
        p = hdr_cells[i].paragraphs[0]
        r = p.add_run(title)
        r.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for idx, (chash, ctype, cdesc) in enumerate(commits):
        row = c_table.rows[idx+1]
        c0, c1, c2 = row.cells[0], row.cells[1], row.cells[2]
        c0.width = Inches(1.2)
        c1.width = Inches(1.8)
        c2.width = Inches(4.0)
        bg = "F8FAFC" if idx % 2 == 0 else "FFFFFF"
        for c in [c0, c1, c2]:
            set_cell_background(c, bg)
            set_cell_margins(c, 60, 60, 100, 100)
        
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(chash)
        r0.font.name = 'Consolas'
        r0.font.size = Pt(8.5)
        r0.bold = True

        p1 = c1.paragraphs[0]
        r1 = p1.add_run(ctype)
        r1.font.size = Pt(8.5)
        r1.font.color.rgb = RGBColor(0x02, 0x84, 0xC7)

        p2 = c2.paragraphs[0]
        r2 = p2.add_run(cdesc)
        r2.font.size = Pt(8.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Section 3: Technical Details
    h3 = doc.add_heading('3. Rincian Teknis Implementasi per Pilar', level=1)
    h3.style.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    details = [
        ("3.1 Pilar A — Keselamatan Data & Storage", [
            ("Proteksi Future Schema", "Menolak pemanggilan reset jika skema storage tergolong 'incompatible' (berasal dari versi aplikasi lebih baru). Menangguhkan pembersihan kunci hingga status terbukti valid/missing. Menyediakan tombol salin dan unduh raw JSON payload pada StorageRecoveryScreen."),
            ("Parser CSV RFC 4180 Mandiri", "Modul lib/csv-parser.ts mengurai baris dengan koma di dalam tanda kutip, escaped double quotes, line breaks, dan validasi UTF-8. Menghilangkan fallback berbahaya '|| 10000'."),
            ("Dedup Engine Impor", "Modul lib/dedup.ts menyaring duplikat transaksi berbasis signature deterministik (ISO date + type + amount + normalized description + category + subcategory + paymentMethod)."),
            ("Pelaporan Rollback Akurat", "Memverifikasi rollback.success secara ketat sebelum melaporkan status; mempertahankan data checkpoint di storage jika terjadi kegagalan tak terduga.")
        ]),
        ("3.2 Pilar B — Navigasi Sublayer LIFO Back-Stack", [
            ("Modul lib/back-stack.ts", "Mengelola tumpukan sublayer (LIFO) dengan pushSublayer, popSublayer, dan dismissSublayer."),
            ("Integrasi Hardware & Browser", "Mendengarkan event hardware back button Android via @capacitor/app App.addListener('backButton') dan browser popstate event."),
            ("Preservasi State Parent", "Menutup sublayer (modal, sheet, dialog) tanpa me-reload aplikasi, menjaga tab aktif, filter tanggal, dan posisi scroll pengguna.")
        ]),
        ("3.3 Pilar C — Standardisasi Live Rupiah Input", [
            ("Modul lib/amount.ts", "Fungsi stripToDigits (menghapus non-digit dan menormalkan leading zero), formatRupiahLive (pemformatan titik ribuan), dan calculateCursorPosition (mencegah kursor melompat saat pemformatan live)."),
            ("Komponen <RupiahInput>", "Prefix 'Rp' terpisah di luar area input sehingga kebal dari penghapusan backspace. Mendukung shortcut natural (50rb, 1,5jt, 100k) dan pembersihan paste terformat."),
            ("Standardisasi Menyeluruh", "Diterapkan pada tab-saku.tsx (anggaran & transfer), edit-transaction-modal.tsx, goal-tracker.tsx (tabungan), dan category-manager.tsx (budget kategori).")
        ]),
        ("3.4 Pilar D — Overhaul UI Pencatatan Transaksi", [
            ("Segmented Type Selector", "3 tombol pill ergonomis (≥44px) dengan diferensiasi warna tegas: Merah untuk Pengeluaran, Hijau untuk Pemasukan, Biru/Cyan untuk Transfer."),
            ("Amount Hero Display", "Desain kartu hero dengan indikator visual dinamis (− / +), label kontekstual, dan quick chip (+10rb, +50rb, +100rb, +500rb)."),
            ("Pengurutan Kategori Dinamis", "Kategori diurutkan berdasarkan histori frekuensi transaksi terbanyak oleh pengguna, serta mengingat kategori terakhir per tipe transaksi."),
            ("Tombol Submit Kontekstual", "Label tombol memuat nominal real-time (cth: 'Catat Pengeluaran Rp 50.000') dengan proteksi double-submit.")
        ]),
        ("3.5 Pilar E — Fitur Baru Jejak Kategori Setahun", [
            ("Analytics Engine lib/stats-category-yearly.ts", "Agregasi 12 bulan penuh kalender (Januari-Desember), pengecualian transfer internal dan simpanan tabungan, penghitungan total tahunan, rata-rata bulanan, porsi persentase, bulan puncak, dan perbandingan vs tahun sebelumnya."),
            ("Komponen Visualisasi category-year-explorer.tsx", "Selector tahun dan kategori horizontal, 5 kartu metrik analitis, grafik batang interaktif Recharts 12 bulan, dan drilldown daftar transaksi saat batang bulan diklik."),
            ("Integrasi Komprehensif", "Dihubungkan pada tab-rekapan-yearly.tsx (mode tahunan) dan tab-rekapan.tsx (rincian kategori tab tren).")
        ]),
        ("3.6 Pilar F — Responsivitas Viewport & Layout", [
            ("Proteksi Overflow 320px+", "Pencegahan horizontal scrolling pada layar kompak melalui css rules max-width: 100vw; overflow-x: hidden;"),
            ("Utilitas Scrollbar Bersih", "Kelas scrollbar-none untuk chip horizontal tanpa scrollbar abu-abu browser."),
            ("Safe Area & Tabular Nums", "Dukungan padding safe-top / safe-bottom serta white-space: nowrap untuk angka moneter.")
        ])
    ]

    for section_title, items in details:
        h_sub = doc.add_heading(section_title, level=2)
        h_sub.style.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
        for it_title, it_desc in items:
            p = doc.add_paragraph()
            r_t = p.add_run(f"• {it_title}: ")
            r_t.bold = True
            r_d = p.add_run(it_desc)
            p.paragraph_format.space_after = Pt(4)

    # Section 4: Hasil Uji
    h4 = doc.add_heading('4. Hasil Uji & Verifikasi Sistem', level=1)
    h4.style.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    test_results = [
        ("Logika Budget Hierarkis", "scripts/test-budget-logic.mjs", "15 Tests", "PASS (100%)"),
        ("SK-003: Storage Corruption Recovery", "scripts/test-storage-corruption.mjs", "18 Tests", "PASS (100%)"),
        ("Future Schema Safety & Recovery", "scripts/test-future-schema-safety.mjs", "11 Tests", "PASS (100%)"),
        ("SK-004: Transactional Restore & CSV", "scripts/test-import-restore.mjs", "16 Tests", "PASS (100%)"),
        ("Sublayer Navigation Back Stack", "scripts/test-navigation-stack.mjs", "7 Tests", "PASS (100%)"),
        ("Live Rupiah Formatting & Parsing", "scripts/test-amount-rupiah.mjs", "4 Groups", "PASS (100%)"),
        ("Jejak Kategori Setahun Analytics", "scripts/test-category-yearly.mjs", "5 Groups", "PASS (100%)"),
        ("SK-001: Scanner Data Personal & Finansial", "scripts/scan-personal-data.mjs", "3 Checks", "PASS (Clean)")
    ]

    t_table = doc.add_table(rows=len(test_results)+1, cols=4)
    t_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_hdr = t_table.rows[0].cells
    t_hdr[0].width = Inches(2.2)
    t_hdr[1].width = Inches(2.3)
    t_hdr[2].width = Inches(1.1)
    t_hdr[3].width = Inches(1.4)
    for i, title in enumerate(["Suite Pengujian", "File Script", "Cakupan", "Hasil Akhir"]):
        set_cell_background(t_hdr[i], "1E293B")
        set_cell_margins(t_hdr[i], 80, 80, 100, 100)
        p = t_hdr[i].paragraphs[0]
        r = p.add_run(title)
        r.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for idx, (sname, sfile, scov, sres) in enumerate(test_results):
        row = t_table.rows[idx+1]
        c0, c1, c2, c3 = row.cells[0], row.cells[1], row.cells[2], row.cells[3]
        c0.width = Inches(2.2)
        c1.width = Inches(2.3)
        c2.width = Inches(1.1)
        c3.width = Inches(1.4)
        bg = "F8FAFC" if idx % 2 == 0 else "FFFFFF"
        for c in [c0, c1, c2, c3]:
            set_cell_background(c, bg)
            set_cell_margins(c, 60, 60, 100, 100)
        
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(sname)
        r0.font.size = Pt(8.5)
        r0.bold = True

        p1 = c1.paragraphs[0]
        r1 = p1.add_run(sfile)
        r1.font.name = 'Consolas'
        r1.font.size = Pt(8)

        p2 = c2.paragraphs[0]
        r2 = p2.add_run(scov)
        r2.font.size = Pt(8.5)

        p3 = c3.paragraphs[0]
        r3 = p3.add_run(sres)
        r3.font.size = Pt(8.5)
        r3.bold = True
        r3.font.color.rgb = RGBColor(0x04, 0x78, 0x57)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Build Verification
    h_bld = doc.add_heading('4.2 Verifikasi Kompilasi & Build', level=2)
    bld_results = [
        ("pnpm exec tsc --noEmit", "Pemeriksaan tipe TypeScript ketat", "0 Errors / Clean"),
        ("pnpm build", "Next.js Web Production Bundle", "Berhasil (3/3 static pages)"),
        ("pnpm build:mobile", "Next.js Mobile Export (BUILD_TARGET=mobile)", "Berhasil (3/3 static pages)"),
        ("git diff --check", "Audit whitespace dan formatting Git", "Bersih (0 warning)")
    ]
    b_table = doc.add_table(rows=len(bld_results)+1, cols=3)
    b_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    b_hdr = b_table.rows[0].cells
    b_hdr[0].width = Inches(2.2)
    b_hdr[1].width = Inches(2.8)
    b_hdr[2].width = Inches(2.0)
    for i, title in enumerate(["Perintah Build", "Deskripsi Target", "Status"]):
        set_cell_background(b_hdr[i], "1E293B")
        set_cell_margins(b_hdr[i], 80, 80, 100, 100)
        p = b_hdr[i].paragraphs[0]
        r = p.add_run(title)
        r.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for idx, (bcmd, bdesc, bres) in enumerate(bld_results):
        row = b_table.rows[idx+1]
        c0, c1, c2 = row.cells[0], row.cells[1], row.cells[2]
        c0.width = Inches(2.2)
        c1.width = Inches(2.8)
        c2.width = Inches(2.0)
        bg = "F8FAFC" if idx % 2 == 0 else "FFFFFF"
        for c in [c0, c1, c2]:
            set_cell_background(c, bg)
            set_cell_margins(c, 60, 60, 100, 100)
        
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(bcmd)
        r0.font.name = 'Consolas'
        r0.font.size = Pt(8.5)
        r0.bold = True

        p1 = c1.paragraphs[0]
        r1 = p1.add_run(bdesc)
        r1.font.size = Pt(8.5)

        p2 = c2.paragraphs[0]
        r2 = p2.add_run(bres)
        r2.font.size = Pt(8.5)
        r2.bold = True
        r2.font.color.rgb = RGBColor(0x04, 0x78, 0x57)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Section 5: Manual Verification Guide
    h5 = doc.add_heading('5. Panduan Verifikasi Pengguna', level=1)
    h5.style.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    p_v = doc.add_paragraph(
        'Pengguna dapat menguji langsung seluruh fungsionalitas aplikasi di lingkungan peramban lokal dengan langkah:\n\n'
        '1. Jalankan server lokal: pnpm dev\n'
        '2. Akses http://localhost:3000 pada peramban.\n'
        '3. Uji Input Rupiah: Buka tab Saku, ubah saldo atau anggaran harian. Perhatikan pemformatan titik ribuan live dan ketik shortcut seperti 50rb atau 1,5jt.\n'
        '4. Uji Modal Catat Transaksi: Tekan tombol +, ganti tipe transaksi (Pengeluaran, Pemasukan, Transfer). Amati segmented pills 44px, hero amount display, dan pengurutan kategori dinamis.\n'
        '5. Uji Navigasi Back Stack: Saat modal terbuka, tekan tombol Back peramban. Modal akan tertutup dengan mulus tanpa mengubah posisi halaman atau me-refresh aplikasi.\n'
        '6. Uji Jejak Kategori Setahun: Buka tab Rekapan, masuk ke mode Tahunan, lalu klik banner "Jejak Kategori Setahun". Amati grafik batang 12 bulan dan klik salah satu bulan untuk melihat drilldown transaksi.'
    )

    # Section 6: Penutup
    h6 = doc.add_heading('6. Kesimpulan & Penutup', level=1)
    h6.style.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    p_c = doc.add_paragraph(
        'Penyelesaian tahap ini memperkokoh integritas data offline SakuKilat sekaligus menyajikan pengalaman pengguna '
        'yang setara dengan aplikasi finansial modern. Seluruh kode telah diuji secara komprehensif dan tersimpan rapi pada '
        'branch feat/navigation-entry-category-responsive siap untuk proses merge selanjutnya.'
    )

    doc.save(output_path)
    print(f"DOCX created: {output_path}")

def generate_html(output_path):
    html_content = """<!DOCTYPE html>
<html lang="id">
<head>
  <meta charset="UTF-8">
  <meta name="viewport" content="width=device-width, initial-scale=1.0">
  <title>Laporan Final Koreksi, UX Upgrade & Category Analytics — SakuKilat</title>
  <style>
    :root {
      --bg: #0B0F19;
      --card: #111827;
      --card-border: #1F2937;
      --text: #F3F4F6;
      --text-muted: #9CA3AF;
      --accent: #38BDF8;
      --accent-green: #10B981;
      --accent-emerald: #059669;
      --accent-amber: #F59E0B;
      --accent-rose: #F43F5E;
      --table-header: #1E293B;
      --code-bg: #1E2433;
    }
    @media print {
      body { background: white !important; color: #1e293b !important; }
      .card { border: 1px solid #cbd5e1 !important; background: white !important; box-shadow: none !important; }
      th { background: #f1f5f9 !important; color: #0f172a !important; }
      td { border-color: #e2e8f0 !important; color: #334155 !important; }
      .badge-green { background: #dcfce7 !important; color: #166534 !important; border-color: #86efac !important; }
      .badge-blue { background: #e0f2fe !important; color: #0369a1 !important; border-color: #7dd3fc !important; }
      .meta-box { background: #f8fafc !important; border-color: #cbd5e1 !important; color: #0f172a !important; }
    }
    * { box-sizing: border-box; }
    body {
      font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, Helvetica, Arial, sans-serif;
      background: var(--bg);
      color: var(--text);
      line-height: 1.6;
      margin: 0;
      padding: 36px 18px;
    }
    .container {
      max-width: 980px;
      margin: 0 auto;
    }
    .header {
      text-align: center;
      margin-bottom: 32px;
      border-bottom: 1px solid var(--card-border);
      padding-bottom: 24px;
    }
    h1 {
      font-size: 26px;
      margin: 0 0 10px 0;
      color: #FFFFFF;
      letter-spacing: -0.02em;
    }
    .subtitle {
      font-size: 14px;
      color: var(--text-muted);
      font-weight: 400;
      max-width: 780px;
      margin: 0 auto;
    }
    .meta-box {
      background: var(--card);
      border: 1px solid var(--card-border);
      border-radius: 12px;
      padding: 18px 24px;
      margin-bottom: 28px;
      display: grid;
      grid-template-columns: repeat(auto-fit, minmax(280px, 1fr));
      gap: 14px;
    }
    .meta-item {
      font-size: 13px;
    }
    .meta-item strong {
      color: var(--text-muted);
      display: block;
      font-size: 11px;
      text-transform: uppercase;
      letter-spacing: 0.05em;
      margin-bottom: 3px;
    }
    .meta-item span {
      color: var(--text);
      font-weight: 600;
    }
    .card {
      background: var(--card);
      border: 1px solid var(--card-border);
      border-radius: 12px;
      padding: 24px;
      margin-bottom: 26px;
      box-shadow: 0 4px 16px rgba(0,0,0,0.25);
    }
    .card-title {
      font-size: 18px;
      font-weight: 700;
      margin-top: 0;
      margin-bottom: 16px;
      color: var(--accent);
      display: flex;
      align-items: center;
      gap: 10px;
    }
    p, li {
      font-size: 14px;
      color: #D1D5DB;
    }
    ul, ol {
      margin: 8px 0;
      padding-left: 24px;
    }
    li {
      margin-bottom: 8px;
    }
    strong {
      color: #FFFFFF;
    }
    table {
      width: 100%;
      border-collapse: collapse;
      margin: 16px 0;
      font-size: 13px;
    }
    th, td {
      padding: 10px 14px;
      border: 1px solid var(--card-border);
      text-align: left;
    }
    th {
      background: var(--table-header);
      color: #FFFFFF;
      font-weight: 600;
    }
    td {
      color: #D1D5DB;
    }
    tr:nth-child(even) td {
      background: rgba(255,255,255,0.02);
    }
    .badge {
      display: inline-block;
      padding: 3px 8px;
      border-radius: 6px;
      font-size: 11px;
      font-weight: 700;
      text-transform: uppercase;
      letter-spacing: 0.03em;
    }
    .badge-green {
      background: rgba(16, 185, 129, 0.15);
      color: #34D399;
      border: 1px solid rgba(16, 185, 129, 0.3);
    }
    .badge-blue {
      background: rgba(56, 189, 248, 0.15);
      color: #7DD3FC;
      border: 1px solid rgba(56, 189, 248, 0.3);
    }
    code {
      font-family: 'SFMono-Regular', Consolas, 'Liberation Mono', Menlo, monospace;
      background: var(--code-bg);
      padding: 2px 6px;
      border-radius: 4px;
      font-size: 12px;
      color: #38BDF8;
    }
    pre {
      background: var(--code-bg);
      border: 1px solid var(--card-border);
      border-radius: 8px;
      padding: 14px;
      overflow-x: auto;
      font-size: 12px;
      line-height: 1.5;
      color: #E2E8F0;
    }
    .grid-2 {
      display: grid;
      grid-template-columns: 1fr 1fr;
      gap: 16px;
    }
    @media (max-width: 720px) {
      .grid-2 { grid-template-columns: 1fr; }
    }
    .stat-card {
      background: rgba(255,255,255,0.03);
      border: 1px solid var(--card-border);
      border-radius: 8px;
      padding: 14px;
    }
    .stat-title {
      font-size: 12px;
      color: var(--text-muted);
      font-weight: 600;
      margin-bottom: 6px;
    }
    .stat-value {
      font-size: 20px;
      font-weight: 700;
      color: #FFFFFF;
    }
  </style>
</head>
<body>
  <div class="container">
    <div class="header">
      <h1>SAKUKILAT — LAPORAN RESMI FINAL CORRECTION, UX UPGRADE & CATEGORY ANALYTICS</h1>
      <div class="subtitle">Dokumentasi Komprehensif Penyelesaian Remediasi Storage, Parser CSV, Navigasi Back Stack, Live Input Rupiah, Desain Form Catat Transaksi, dan Modul Jejak Kategori Setahun</div>
    </div>

    <div class="meta-box">
      <div class="meta-item">
        <strong>Tanggal Eksekusi</strong>
        <span>11 September 2026</span>
      </div>
      <div class="meta-item">
        <strong>Branch Kerja Aktif</strong>
        <span><code>feat/navigation-entry-category-responsive</code></span>
      </div>
      <div class="meta-item">
        <strong>Base Commit & Remote</strong>
        <span><code>f01f76e</code> (origin/main)</span>
      </div>
      <div class="meta-item">
        <strong>Current HEAD Local</strong>
        <span><code>1bae659</code> (8 commits ahead)</span>
      </div>
      <div class="meta-item">
        <strong>Status Pelaksanaan</strong>
        <span style="color: #34D399;">✓ SELESAI PENUH (8/8 Commits, All Tests Green)</span>
      </div>
      <div class="meta-item">
        <strong>Repository</strong>
        <span>C:\\Users\\HYPE AMD\\Projects\\SakuKilat</span>
      </div>
    </div>

    <!-- Card 1 -->
    <div class="card">
      <div class="card-title">1. Ringkasan Eksekutif</div>
      <p>Sesuai instruksi dan panduan rekayasa perangkat lunak SakuKilat, seluruh pekerjaan final remediation dan peningkatan kapabilitas aplikasi telah diselesaikan 100% tanpa menambahkan library pihak ketiga di luar Recharts yang sudah ada, tanpa memperkenalkan fitur yang tidak diminta (seperti rekonsiliasi bank otomatis atau sinkronisasi cloud), serta dengan integritas data lokal yang terjaga penuh.</p>
      
      <div class="grid-2">
        <div class="stat-card">
          <div class="stat-title">REGRESSION TEST SUITES</div>
          <div class="stat-value" style="color: #34D399;">7 / 7 LULUS (79 Tests)</div>
        </div>
        <div class="stat-card">
          <div class="stat-title">SECURITY SCAN SUITES</div>
          <div class="stat-value" style="color: #38BDF8;">1 / 1 CLEAN (3 Checks)</div>
        </div>
        <div class="stat-card">
          <div class="stat-title">NEXT.JS WEB BUILD</div>
          <div class="stat-value" style="color: #34D399;">COMPILED (0 Error)</div>
        </div>
        <div class="stat-card">
          <div class="stat-title">MOBILE EXPORT BUILD</div>
          <div class="stat-value" style="color: #34D399;">SUCCESS (3 Pages Static)</div>
        </div>
      </div>
    </div>

    <!-- Card 2 -->
    <div class="card">
      <div class="card-title">2. Histori Commit di Branch Kerja</div>
      <p>Pekerjaan dilakukan secara disiplin dalam commit-commit atomik yang masing-masing disertai unit test dan verifikasi.</p>
      <table>
        <thead>
          <tr>
            <th style="width: 15%;">Hash</th>
            <th style="width: 25%;">Scope & Type</th>
            <th style="width: 60%;">Pesan & Deskripsi Perubahan</th>
          </tr>
        </thead>
        <tbody>
          <tr>
            <td><code>4e1eed3</code></td>
            <td><span class="badge badge-blue">fix(storage)</span></td>
            <td>Prevent destructive incompatible reset and preserve future keys</td>
          </tr>
          <tr>
            <td><code>38bfc6d</code></td>
            <td><span class="badge badge-blue">fix(import)</span></td>
            <td>Strict CSV parser, dedup engine, and honest rollback reporting</td>
          </tr>
          <tr>
            <td><code>2497937</code></td>
            <td><span class="badge badge-blue">test(navigation)</span></td>
            <td>Define nested back-stack behavior</td>
          </tr>
          <tr>
            <td><code>3c336bd</code></td>
            <td><span class="badge badge-blue">fix(navigation)</span></td>
            <td>Return sublayers to their parent via LIFO back-stack</td>
          </tr>
          <tr>
            <td><code>e0546a0</code></td>
            <td><span class="badge badge-blue">fix(amount)</span></td>
            <td>Apply live Rupiah formatting consistently across forms</td>
          </tr>
          <tr>
            <td><code>57e9bf7</code></td>
            <td><span class="badge badge-blue">feat(entry)</span></td>
            <td>Improve expense and income entry UI</td>
          </tr>
          <tr>
            <td><code>fae2f13</code></td>
            <td><span class="badge badge-blue">feat(category-year)</span></td>
            <td>Add yearly category explorer</td>
          </tr>
          <tr>
            <td><code>1bae659</code></td>
            <td><span class="badge badge-blue">fix(responsive)</span></td>
            <td>Optimize supported viewport sizes</td>
          </tr>
        </tbody>
      </table>
    </div>

    <!-- Card 3 -->
    <div class="card">
      <div class="card-title">3. Detail Remediasi dan Fitur per Bagian</div>

      <h3>3.1 Bagian A — Keselamatan Data & Storage</h3>
      <ul>
        <li><strong>Perlindungan Skema Masa Depan:</strong> <code>lib/storage.ts</code> memblokir eksekusi reset pada status <code>incompatible</code>. Unknown keys dari versi skema lebih baru tidak dihapus. UI karantina <code>StorageRecoveryScreen</code> kini menyembunyikan tombol reset destruktif dan memberikan tombol salin/unduh raw JSON.</li>
        <li><strong>Parser CSV RFC 4180 Ketat:</strong> Modul mandiri <code>lib/csv-parser.ts</code> mendukung koma di dalam tanda kutip, escaped quotes, multiline breaks, dan validasi UTF-8. Fallback berbahaya <code>|| 10000</code> dihapus total.</li>
        <li><strong>Dedup Engine Transaksi Impor:</strong> Modul <code>lib/dedup.ts</code> menyaring transaksi ganda berdasarkan signature deterministik (ISO date + type + amount + normalized description + category + subcategory + paymentMethod).</li>
        <li><strong>Pelaporan Rollback yang Jujur:</strong> Sistem memvalidasi status riil eksekusi rollback sebelum menampilkan notifikasi keberhasilan semu.</li>
      </ul>

      <h3>3.2 Bagian B — Navigasi Sublayer dan Back Stack</h3>
      <ul>
        <li><strong>Modul Navigasi LIFO:</strong> <code>lib/back-stack.ts</code> melacak hirarki modal dan sheet yang terbuka.</li>
        <li><strong>Hardware Back & Popstate:</strong> Mengaitkan event hardware back button Capacitor dan browser popstate ke stack LIFO sehingga setiap penekanan tombol kembali menutup sublayer teratas secara berurutan.</li>
        <li><strong>Isolasi State Parent:</strong> Menutup modal tidak mereset state tab induk, nilai filter tanggal, ataupun posisi scroll pengguna.</li>
      </ul>

      <h3>3.3 Bagian C — Live Rupiah Input</h3>
      <ul>
        <li><strong>Modul Pendukung:</strong> <code>lib/amount.ts</code> menyediakan <code>stripToDigits</code>, <code>formatRupiahLive</code>, dan kalkulasi posisi kursor <code>calculateCursorPosition</code>.</li>
        <li><strong>Komponen <code>&lt;RupiahInput&gt;</code>:</strong> Prefix "Rp" terisolasi di luar input teks (tidak dapat terhapus backspace), pemisah ribuan otomatis tanpa loncatan kursor, dan normalisasi paste/shortcuts natural (<code>50rb</code>, <code>1,5jt</code>).</li>
        <li><strong>Penggantian Menyeluruh:</strong> Digunakan pada <code>tab-saku.tsx</code>, <code>edit-transaction-modal.tsx</code>, <code>goal-tracker.tsx</code>, dan <code>category-manager.tsx</code>.</li>
      </ul>

      <h3>3.4 Bagian D — UI Pencatatan Transaksi Baru</h3>
      <ul>
        <li><strong>Segmented Type Selector:</strong> Tiga tombol segmented lega (tinggi 44px) dengan diferensiasi warna kontras (Pengeluaran merah, Pemasukan hijau, Transfer biru).</li>
        <li><strong>Amount Hero Section:</strong> Area input nominal visual dominan dengan badge <code>−</code> atau <code>+</code>, label dinamis, dan quick chip (+10rb, +50rb, +100rb, +500rb).</li>
        <li><strong>Pengurutan Kategori Dinamis:</strong> Kategori disusun berdasarkan riwayat frekuensi pemakaian pengguna; mengingat kategori terakhir per tipe.</li>
        <li><strong>Tombol Simpan Adaptif:</strong> Menampilkan nominal real-time (cth: <em>Catat Pengeluaran Rp 75.000</em>) serta pencegahan double-submit.</li>
      </ul>

      <h3>3.5 Bagian E — Fitur Baru: Jejak Kategori Setahun</h3>
      <ul>
        <li><strong>Statistik Engine:</strong> <code>lib/stats-category-yearly.ts</code> mengagregasi 12 bulan kalender secara akurat dengan mengecualikan transfer internal dan simpanan tabungan.</li>
        <li><strong>Visualisasi Explorer:</strong> <code>components/category-year-explorer.tsx</code> menyediakan selector tahun/kategori, 5 kartu metrik analitis, grafik batang Recharts 12 bulan interaktif, dan drilldown transaksi bulanan saat batang diklik.</li>
        <li><strong>Integrasi:</strong> Dapat diakses langsung dari <code>tab-rekapan-yearly.tsx</code> dan <code>tab-rekapan.tsx</code>.</li>
      </ul>

      <h3>3.6 Bagian F — Responsivitas Viewport</h3>
      <ul>
        <li>Proteksi overflow horizontal pada layar sempit (320px+) melalui <code>max-width: 100vw; overflow-x: hidden;</code> pada <code>app/globals.css</code>.</li>
        <li>Utilitas <code>.scrollbar-none</code> untuk geser chip kategori horizontal tanpa gangguan scrollbar browser.</li>
        <li>Penyematan safe-area top dan bottom untuk kompatibilitas perangkat mobile modern.</li>
      </ul>
    </div>

    <!-- Card 4 -->
    <div class="card">
      <div class="card-title">4. Hasil Uji Otomatis & Verifikasi Build</div>
      
      <table>
        <thead>
          <tr>
            <th>Suite Pengujian</th>
            <th>File Script</th>
            <th>Cakupan</th>
            <th>Hasil</th>
          </tr>
        </thead>
        <tbody>
          <tr>
            <td>Logika Budget Hierarkis</td>
            <td><code>scripts/test-budget-logic.mjs</code></td>
            <td>15 Tests</td>
            <td><span class="badge badge-green">PASS (100%)</span></td>
          </tr>
          <tr>
            <td>SK-003: Storage Corruption Recovery</td>
            <td><code>scripts/test-storage-corruption.mjs</code></td>
            <td>18 Tests</td>
            <td><span class="badge badge-green">PASS (100%)</span></td>
          </tr>
          <tr>
            <td>Future Schema Safety & Recovery</td>
            <td><code>scripts/test-future-schema-safety.mjs</code></td>
            <td>11 Tests</td>
            <td><span class="badge badge-green">PASS (100%)</span></td>
          </tr>
          <tr>
            <td>SK-004: Transactional Restore & CSV</td>
            <td><code>scripts/test-import-restore.mjs</code></td>
            <td>16 Tests</td>
            <td><span class="badge badge-green">PASS (100%)</span></td>
          </tr>
          <tr>
            <td>Sublayer Navigation Back Stack</td>
            <td><code>scripts/test-navigation-stack.mjs</code></td>
            <td>7 Tests</td>
            <td><span class="badge badge-green">PASS (100%)</span></td>
          </tr>
          <tr>
            <td>Live Rupiah Formatting & Parsing</td>
            <td><code>scripts/test-amount-rupiah.mjs</code></td>
            <td>4 Groups</td>
            <td><span class="badge badge-green">PASS (100%)</span></td>
          </tr>
          <tr>
            <td>Jejak Kategori Setahun Analytics</td>
            <td><code>scripts/test-category-yearly.mjs</code></td>
            <td>5 Groups</td>
            <td><span class="badge badge-green">PASS (100%)</span></td>
          </tr>
          <tr>
            <td>SK-001: Scanner Data Personal & Finansial</td>
            <td><code>scripts/scan-personal-data.mjs</code></td>
            <td>3 Checks</td>
            <td><span class="badge badge-green">CLEAN (0 Leaks)</span></td>
          </tr>
        </tbody>
      </table>

      <h3>Verifikasi Build Produksi</h3>
      <table>
        <thead>
          <tr>
            <th>Perintah</th>
            <th>Deskripsi Target</th>
            <th>Hasil</th>
          </tr>
        </thead>
        <tbody>
          <tr>
            <td><code>pnpm exec tsc --noEmit</code></td>
            <td>Pemeriksaan Strict Tipe TypeScript</td>
            <td><span class="badge badge-green">0 Errors / Lulus</span></td>
          </tr>
          <tr>
            <td><code>pnpm build</code></td>
            <td>Next.js Web Production Build</td>
            <td><span class="badge badge-green">Compiled Successfully (3/3 pages)</span></td>
          </tr>
          <tr>
            <td><code>pnpm build:mobile</code></td>
            <td>Next.js Mobile Export (BUILD_TARGET=mobile)</td>
            <td><span class="badge badge-green">Compiled Successfully (3/3 pages)</span></td>
          </tr>
          <tr>
            <td><code>git diff --check</code></td>
            <td>Audit Whitespace & Format Git</td>
            <td><span class="badge badge-green">Clean (0 warnings)</span></td>
          </tr>
        </tbody>
      </table>
    </div>

    <!-- Card 5 -->
    <div class="card">
      <div class="card-title">5. Panduan Pengujian Langsung oleh Pengguna</div>
      <p>Untuk menguji secara visual melalui peramban:</p>
      <ol>
        <li>Jalankan dev server dengan perintah: <code>pnpm dev</code></li>
        <li>Buka <code>http://localhost:3000</code> di browser Anda.</li>
        <li><strong>Input Rupiah:</strong> Buka tab <strong>Saku</strong>, ubah saldo atau limit harian. Coba ketik angka atau shortcut <code>50rb</code> / <code>1,5jt</code>.</li>
        <li><strong>Form Catat:</strong> Klik tombol <strong>+</strong>. Pilih Pengeluaran / Pemasukan / Transfer. Rasakan hero nominal dan perhatikan kategori teratas berdasarkan frekuensi.</li>
        <li><strong>Tombol Kembali (Back):</strong> Saat modal terbuka, tekan tombol kembali pada peramban. Modal akan tertutup dengan mulus tanpa me-reload aplikasi.</li>
        <li><strong>Jejak Kategori Setahun:</strong> Masuk tab <strong>Rekapan</strong> ➔ <strong>Tahunan</strong> ➔ Klik <strong>Jejak Kategori Setahun</strong>. Amati visualisasi batang 12 bulan dan klik salah satu batang untuk melihat daftar transaksi.</li>
      </ol>
    </div>

    <div class="card" style="text-align: center; border-color: rgba(16, 185, 129, 0.3);">
      <p style="margin: 0; color: #34D399; font-weight: 600;">Repository dalam kondisi bersih (*clean working tree*). Seluruh perubahan siap ditinjau dan digabungkan (*merge*).</p>
    </div>
  </div>
</body>
</html>
"""
    with open(output_path, "w", encoding="utf-8") as f:
        f.write(html_content)
    print(f"HTML created: {output_path}")

if __name__ == "__main__":
    docs_dir = os.path.join(os.path.dirname(__file__), "..", "docs")
    docx_path = os.path.join(docs_dir, "LAPORAN_FINAL_KOREKSI_UX_SAKUKILAT.docx")
    html_path = os.path.join(docs_dir, "LAPORAN_FINAL_KOREKSI_UX_SAKUKILAT.html")
    
    generate_docx(docx_path)
    generate_html(html_path)
